"""DoD RFO class deviations (acq.osd.mil DPAP page, fetched by hand because the host's DoD certificate is not trusted).

One row per deviation memo: tracking number, effective date, and the DFARS 252 clause numbers its attachment 1 mentions.
Files live under data/raw/manual/dod/files; manifest.json there records source, sizes and hashes.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from fedproc_ledger.registry.deviations import iso_date

_TRACK = re.compile(r"(?:DARS\s+)?Tracking\s+Number:\s*(\d{4}-O\d{4})", re.I)
_EFFECTIVE = re.compile(r"Effective\s+([A-Z][a-z]+\.?\s+\d{1,2},\s+20\d{2})")
_SUPERSEDES = re.compile(
    r"supersedes\s+Class\s+Deviation\s+[\d\-\s]*O\s?\d{4}(?:,?\s+Revision\s+\d+)?,?\s+issued\s+on\s+([A-Z][a-z]+\.?\s+\d{1,2},\s+20\d{2})"
)
_CLAUSE_252 = re.compile(r"(?<!\d)252\.\d{3}-\d{4}(?!\d)")
BASE = "https://www.acq.osd.mil/dpap/dars/classdev/DFARS_RFO/"


def parse_memo(text: str) -> dict[str, Any]:
    """Revisions say "Effective immediately ... supersedes ... issued on <date>": no effective date is stated, so
    `date` stays None and the superseded issue date is kept separately."""
    flat = re.sub(r"\s+", " ", text)
    track = _TRACK.search(flat)
    eff = _EFFECTIVE.search(flat)
    sup = _SUPERSEDES.search(flat)
    return {
        "deviation_number": track.group(1) if track else None,
        "date": iso_date(eff.group(1)) if eff else None,
        "revision": sup is not None,
        "supersedes_issued": iso_date(sup.group(1)) if sup else None,
        "immediate": "Effective immediately" in flat,
    }


def clauses_252(text: str) -> list[str]:
    return sorted(set(_CLAUSE_252.findall(text)))


def build_rows(root: Path) -> list[dict[str, Any]]:
    import docx
    import pymupdf

    rows = []
    for memo in sorted(set(root.glob("*/*TAB_A_*.pdf"))):
        part = memo.parent.name
        pdf: Any = pymupdf.open(memo)  # type: ignore[no-untyped-call]
        head = parse_memo("\n".join(p.get_text() for p in pdf))
        nums: list[str] = []
        for att in memo.parent.glob("*TAB_A1_*.docx"):
            d: Any = docx.Document(str(att))
            cells = (c.text for t in d.tables for r in t.rows for c in r.cells)
            nums = clauses_252("\n".join([p.text for p in d.paragraphs] + list(cells)))
        rows.append(
            {
                "kind": "agency",
                "agency": "DoD",
                "deviation_number": head["deviation_number"],
                "date": head["date"],
                "date_updated": None,
                "revision": head["revision"],
                "supersedes_issued": head["supersedes_issued"],
                "clause_numbers_mentioned": nums,
                "url": BASE + str(memo.relative_to(root)).replace(" ", "%20"),
                "parsed": head["date"] is not None or head["immediate"],
                "header": part,
            }
        )
    return rows
