import docx
from lxml import etree

from fedproc_ledger.extract.docx import W14, W, extract_docx, paragraph_line
from fedproc_ledger.extract.pdf import CHECKED_M, UNCHECKED_M, UNKNOWN_M


def para(xml_body: str):
    return etree.fromstring(f'<w:p xmlns:w="{W}" xmlns:w14="{W14}">{xml_body}</w:p>')


def run(text):
    return f"<w:r><w:t xml:space='preserve'>{text}</w:t></w:r>"


def sdt_checkbox(checked: bool, glyph: str):
    return (
        f'<w:sdt><w:sdtPr><w14:checkbox><w14:checked w14:val="{1 if checked else 0}"/></w14:checkbox></w:sdtPr>'
        f"<w:sdtContent><w:r><w:t>{glyph}</w:t></w:r></w:sdtContent></w:sdt>"
    )


def legacy_checkbox(checked: bool):
    chk = '<w:checked w:val="1"/>' if checked else ""
    return f'<w:r><w:fldChar w:fldCharType="begin"><w:ffData><w:checkBox><w:default w:val="0"/>{chk}</w:checkBox></w:ffData></w:fldChar></w:r>'


def test_content_control_checkbox_state_comes_from_the_control_not_the_glyph():
    ln = paragraph_line(para(sdt_checkbox(True, "☒") + run(" FAR 52.243-5 Changes")))
    assert ln.text_layout.startswith(f"{CHECKED_M} FAR 52.243-5") and [b.source for b in ln.boxes] == ["widget"]
    ln2 = paragraph_line(para(sdt_checkbox(False, "☐") + run(" FAR 52.243-4")))
    assert ln2.text_layout.startswith(f"{UNCHECKED_M} FAR 52.243-4") and len(ln2.boxes) == 1  # not counted twice
    assert "☒" in ln.text_plain  # the plain view keeps the control's inner glyph, like a simple extractor would


def test_legacy_form_checkbox():
    assert paragraph_line(para(legacy_checkbox(True) + run(" 52.203-6"))).text_layout.startswith(CHECKED_M)
    assert paragraph_line(para(legacy_checkbox(False) + run(" 52.203-6"))).text_layout.startswith(UNCHECKED_M)


def test_glyph_boxes_in_runs_and_symbol_runs_of_unknown_meaning():
    ln = paragraph_line(para(run("☒ 52.212-4 ") + run("then ☐ 52.212-5")))
    assert ln.text_layout == f"{CHECKED_M} 52.212-4 then {UNCHECKED_M} 52.212-5"
    sym = paragraph_line(para('<w:r><w:sym w:font="Wingdings" w:char="F0FE"/></w:r>' + run(" 52.204-21")))
    assert sym.text_layout.startswith(UNKNOWN_M) and sym.boxes[0].state == "?" and "F0FE" in sym.boxes[0].detail


def test_a_tick_counts_only_at_the_start_of_the_paragraph():
    start = paragraph_line(para(run("✔ 52.203-6 Restrictions")))
    later = paragraph_line(para(run("Selected ") + run("✔ in the middle")))
    assert start.text_layout.startswith(CHECKED_M) and CHECKED_M not in later.text_layout and "✔" in later.text_layout


def test_whole_document_with_a_table(tmp_path):
    d = docx.Document()
    d.add_paragraph("Contract Terms and Conditions Required to Implement Statutes")
    d.add_paragraph("☒ (1) 52.203-6 Restrictions on Subcontractor Sales")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "☐"
    t.cell(0, 1).text = "52.204-10 Reporting Executive Compensation"
    p = tmp_path / "x.docx"
    d.save(p)
    (page,) = extract_docx(p)
    assert page.page == 1 and page.scanned is False
    assert f"{CHECKED_M} (1) 52.203-6" in page.text_layout
    assert f"{UNCHECKED_M} | 52.204-10" in page.text_layout
    assert page.box_counts == {"glyph:X": 1, "glyph: ": 1}
